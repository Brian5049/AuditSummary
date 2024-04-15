import os
import csv
import time
import getpass
import openpyxl
import subprocess
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import Image
from reportlab.lib.units import inch
from reportlab.lib.utils import simpleSplit
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics.shapes import Drawing
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.graphics.charts.piecharts import Pie

folder = input("What folder would you like to summarize? ")

# Variables:
# These will be used throughout the process and are necessary to the script
userName = getpass.getuser()
csvPath = f"C:/Users/{userName}/Desktop/AuditSummaries/{folder}-summary.csv"
csvFile = f"{folder}-summary.csv"
# Please change your colors. I'll leave these here, but it would be nice
# if you could come up with your own scheme. Maybe something to match
# your logo (;-)
OxfordBlue = colors.HexColor('#0f243e')
MoonstoneBlue = colors.HexColor('#60ACC3')
LimeGreen = colors.HexColor('#00FF00')
Red = colors.HexColor('#FF0000')

# Set the folder path. This file path in my instance included a SharePoint drive.
# I know it works but you will need to play with this between the three scripts 
# on your end to debug and verify
folder_path = rf"C:\Users\{userName}\{folder}"
# Initialize a variable to count the spreadsheets with non-null values
count = 0
# List all files in the directory
files = os.listdir(folder_path)
# Iterate over each file to get not_null_counts. These were specific
# to my use case. You may want to change this to reflect what you 
# are trying to do.
for file_name in files:
    if file_name.endswith('.xlsx'):
        file_path = os.path.join(folder_path, file_name)
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        for row in range(17, 35):
            cell_value = ws.cell(row=row, column=8).value
            if cell_value is not None:
                count += 1
                break
FailedDevices = int(count) # Another thing that you might not need
# File existence check and delete
def DeleteFile(PDFpath):
    if os.path.exists(PDFpath):
        os.remove(PDFpath)
        print(f"File '{PDFpath}' has been deleted.")
    else:
        print(f"File '{PDFpath}' does not exist.")
PDFpath = csvPath

# PowerShell scripts
powershell_script1 = "./process1.ps1"
powershell_script2 = "./process2.ps1"
subprocess.run(["powershell.exe", "-File", powershell_script1, folder], shell=True)
DeleteFile(PDFpath)
subprocess.run(["powershell.exe", "-File", powershell_script2, folder], shell=True)

with open(csvPath, 'r') as file:
    total_devices_audited = file.readline().strip().split(',')[1]
    # Remove quotation marks
    total_devices_audited = total_devices_audited.replace('"', '')
    # Convert the string to an integer
    DevicesAudited = int(total_devices_audited)

with open(csvPath, 'r') as file:
    questions = []
    yes_counts = []
    no_counts = []
    not_null_counts = []
    reader = csv.reader(file)
    next(reader)
    # After headers are skipped, you will need to alter this
    # code to fit your use case
    headers = next(reader)
    question_index = headers.index('Audit Questions')
    yes_count_index = headers.index('YesCount')
    no_count_index = headers.index('NoCount')
    not_null_count_index = headers.index('NotNullCount')

    for row in reader:
        # Again, check all of this against your goals. Change
        # variables and how you do counts if needed. This is
        # only the way I needed it.
        questions.append(row[question_index])
        yes_counts.append(int(row[yes_count_index]))
        no_counts.append(int(row[no_count_index]))
        not_null_counts.append(int(row[not_null_count_index]))
        EditsMade = sum(not_null_counts)

def CreatePDF():
    doc = rf"C:\Users\{userName}\Desktop\AuditSummaries\{folder}.pdf"

    # Don't forget to change this code!!!
    $MyLogo = "Whatever logo you may need for your pdf"
    
    # Static variables
    logo_path = 'D:\Docs\Work\AuditScript\AuditSummaryScript\$MyLogo'
    logo = Image(logo_path)
    
#========================================================================#
# Now I'm telling you, this is where it's going to get a little crazy... #   
#========================================================================#
    
    # We are creating out final document here
    c = canvas.Canvas(doc, pagesize=letter)
    
    def header(canvas, c):
        c.setStrokeColor(OxfordBlue)
        c.setFont("Times-Roman", 10)
        c.drawString(0.5*inch, letter[1] - 0.25 * inch, f"WHT 4Advisors CyberSecurity Audit")
        c.drawRightString(letter[0] - 0.5*inch, letter[1] - 0.25 * inch, folder)
        c.setLineWidth(2)
        c.line(0.5*inch, letter[1] - 0.4 * inch, letter[0] - 0.5*inch, letter[1] - 0.4 * inch)

    def footer(canvas, c):
        c.saveState()
        c.setStrokeColor(OxfordBlue)
        c.setLineWidth(2)
        c.line(0.5*inch, 0.4*inch, letter[0] - 0.5*inch, 0.4*inch)
        c.setFont("Times-Roman", 10)
        c.drawString(0.5*inch, 0.15*inch, f"Page {c.getPageNumber()}")
        c.drawRightString(letter[0] - 0.5*inch, 0.15*inch, "wht4advisors.com")

    # We are setting the objects for the first page only. We will do all the other pages later
    def firstPage(canvas, c):
        # You will most likely have to adjust this as your logo will not be the same as mine
        c.drawImage(logo_path, 0.5*inch, 8.4*inch, width=3.25*inch, height=2*inch)
        c.setFillColor(OxfordBlue)
        c.setFont('Times-Roman', 26)
        nameText = f"{folder}:"
        x = 50
        y = 550
        c.drawString(x, y, nameText)
        # Don't forget to change the TITLE!!!
        titleText = "Whatever silly report summary I'm working on!"
        x = 300
        y = 615  
        width = 200
        current_y = y
        lines = titleText.split('\n')
        for line in lines:
            words = line.split(' ')
            current_line = ''
            for word in words:
                if stringWidth(current_line + ' ' + word, 'Times-Roman', 12) < width:
                    current_line += ' ' + word
                else:
                    c.drawString(x, current_y, current_line)
                    current_y -= 25
                    current_line = word
            c.drawString(x, current_y, current_line)
        # Add statement paragraph
        c.setFont('Times-Roman', 12)

        # Don't forget to change this!!!
        paraText = "I advise you to add this text as early as you can. It will help you get a feel for how much space it's going to take up. Lorem ipsum will work, but I think you should just go ahead and type it out... There is another paragraph further down..."
        
        x = 50
        y = 500
        width = 500
        current_y = y
        lines = paraText.split('\n')
        for line in lines:
            words = line.split(' ')
            current_line = ''
            for word in words:
                if stringWidth(current_line + ' ' + word, 'Times-Roman', 12) < width:
                    current_line += ' ' + word
                else:
                    c.drawString(x, current_y, current_line)
                    current_y -= 15
                    current_line = word
            c.drawString(x, current_y, current_line)
            current_y -= 15
        # This pie chart will only work if you do your variables correctly
        failed_percentage = (FailedDevices / DevicesAudited) * 100
        passed_percentage = ((DevicesAudited - FailedDevices) / DevicesAudited) * 100
        pie = Pie()
        pie.x = 40
        pie.y = 100
        pie.width = 200
        pie.height = 200
        pie.data = [failed_percentage, passed_percentage]
        pie.slices.strokeWidth = 0.25
        pie.slices[0].fillColor = OxfordBlue
        pie.slices[1].fillColor = MoonstoneBlue
        pie.slices[1].popout = 5
        pie.sideLabels = 1
        pie.simpleLabels = 0
        pie.slices.strokeColor = OxfordBlue
        pie_chart_drawing = Drawing(width=pie.width, height=pie.height)
        pie_chart_drawing.add(pie)
        pie_chart_drawing.drawOn(c, 40, 100)
        c.setFont("Times-Roman", 16)
        c.setFillColor(OxfordBlue)
        # Remember to check these variables, you may have changed them a couple of days ago.
        c.drawString(300, 320, f"Failed Devices: {FailedDevices} ({failed_percentage:.2f})%")
        c.drawString(300, 300, f"Devices Audited: {DevicesAudited}")
        # Add statement paragraph
        c.setFont('Times-Roman', 12)
        statement = f"""
        I told you there was going to be another paragraph! lol You didn't believe me... Write this one out as soon as you can as well. You really have to work with the formatting of this stuff...
        """
        x = 300
        y = 250
        width = 250
        current_y = y
        lines = statement.split('\n')
        for line in lines:
            words = line.split(' ')
            current_line = ''
            for word in words:
                if stringWidth(current_line + ' ' + word, 'Times-Roman', 12) < width:
                    current_line += ' ' + word
                else:
                    c.drawString(x, current_y, current_line)
                    current_y -= 15
                    current_line = word
            c.drawString(x, current_y, current_line)
            current_y -= 15
        footer(c, c)
    
    # Now we get to format all other pages!!! Yay!!!
    def otherPages(canvas, c):
        max_bar_length = 400
        c.setFont("Times-Roman", 12)
        background_color = MoonstoneBlue
        filled_color = OxfordBlue
        x_pos = 100
        bar_spacing = 40
        questions = []
        not_null_counts = []
        top_margin = 1
        bottom_margin = 1
        available_height = letter[1] - top_margin * inch - bottom_margin * inch
        # This was a little tricky but it should check for space before it just throws a 
        # a bar graph in there. You'll still probably have to play with it though.
        if available_height > bar_spacing:
            with open(csvPath, 'r') as file:
                reader = csv.reader(file)
                next(reader)
                headers = next(reader)
                
                # Please remember your variables...
                question_index = headers.index('Audit Questions')
                not_null_count_index = headers.index('NotNullCount')
                rows = sorted(reader, key=lambda row: int(row[not_null_count_index]), reverse=True)
                
                # Calculate the starting y position for the loop
                start_y_pos = letter[1] - top_margin * inch
                for row in rows:
                    question = row[question_index]
                    not_null_count = int(row[not_null_count_index])
                    questions.append(question)
                    not_null_counts.append(not_null_count)
                    # Calculate the remaining space after subtracting the content height from the available height
                    remaining_space = start_y_pos - bottom_margin * inch
                    if remaining_space > bar_spacing:
                        # Split question text into multiple lines if it exceeds the maximum width. These are some
                        # funky hoops I had to jump through with my project. You might not have to do all this. 
                        # Don't forget to debug as you go LOL
                        lines = simpleSplit(question, "Times-Roman", 12, max_bar_length)
                        line_height = 12
                        for line in lines:
                            c.setFillColor(filled_color)
                            c.drawString(x_pos, start_y_pos, line)
                            start_y_pos -= line_height
                        # This is going to be all about personal preference and conformity to your data set
                        filled_width = max_bar_length * (not_null_count / DevicesAudited)
                        bar_graph_y_pos = start_y_pos - 5
                        c.setStrokeColor(background_color)
                        c.setFillColor(background_color)
                        c.rect(x_pos, bar_graph_y_pos, max_bar_length, 10, fill=1)
                        c.setStrokeColor(filled_color)
                        c.setFillColor(filled_color)
                        c.rect(x_pos, bar_graph_y_pos, filled_width, 10, fill=1)
                        # Add the number associated with the bar graph inside the bar
                        if not_null_count > 0:
                            number_text = f"{not_null_count} Devices Failed"
                            text_width = c.stringWidth(number_text, "Times-Roman", 12)
                            text_x = x_pos + (max_bar_length - text_width) / 2  # Center horizontally <Super Helpful
                            text_y = bar_graph_y_pos
                            c.setFillColor(Red)
                            c.drawString(text_x, text_y, number_text)
                        else:
                            number_text = f"All Passed!"
                            text_width = c.stringWidth(number_text, "Times-Roman", 12)
                            text_x = x_pos + (max_bar_length - text_width) / 2  # Center horizontally
                            text_y = bar_graph_y_pos
                            c.setFillColor(LimeGreen)
                            c.drawString(text_x, text_y, number_text)
                        # Update start_y_pos for the next question
                        start_y_pos -= bar_spacing
                    else:
                        c.showPage()
                        # Add header for every new page
                        header(canvas, c)
                        start_y_pos = letter[1] - top_margin * inch
        else:
            # If there's not enough space to fit even one bar graph, print a message or handle accordingly
            print("Not enough space to fit even one bar graph on the page.")

    firstPage(c, c)
    c.showPage()
    header(c, c)
    c.setFillColor(OxfordBlue)
    footer(c, c)
    otherPages(c, c)
    c.setFillColor(OxfordBlue)
    footer(c, c)
    c.save()

CreatePDF()
DeleteFile(PDFpath)
